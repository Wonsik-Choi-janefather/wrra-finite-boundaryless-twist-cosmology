from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper"


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])


def set_font(run, western, east_asia, size=None, bold=None, color="000000"):
    run.font.name = western
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_by_name(styles, name):
    for style in styles:
        if style.name == name:
            return style
    raise KeyError(name)


def style_doc(path: Path, korean: bool):
    doc = Document(path)
    # LibreOffice's DOCX importer can prefer w:ascii/w:hAnsi over w:eastAsia
    # even for Hangul.  Use the installed CJK family in every font slot for
    # the Korean edition so the PDF contains visible, embedded Korean glyphs.
    body_west = "Noto Sans KR" if korean else "Noto Serif"
    body_east = "Noto Sans KR" if korean else "Noto Serif"
    head_west = "Noto Sans KR" if korean else "Noto Sans"
    head_east = "Noto Sans KR" if korean else "Noto Sans"

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = style_by_name(styles, "Normal")
    normal.font.name = body_west
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), body_east)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.16
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True

    title = style_by_name(styles, "Title")
    title.font.name = head_west
    title._element.rPr.rFonts.set(qn("w:eastAsia"), head_east)
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(90)
    title.paragraph_format.space_after = Pt(16)

    if "Subtitle" in styles:
        subtitle = style_by_name(styles, "Subtitle")
        subtitle.font.name = head_west
        subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), head_east)
        subtitle.font.size = Pt(14)
        subtitle.font.italic = False
        subtitle.font.color.rgb = RGBColor(0, 0, 0)
        subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.space_after = Pt(26)

    for style_name, size, before, after in (
        ("Heading 1", 16, 14, 7),
        ("Heading 2", 13, 11, 5),
        ("Heading 3", 11, 9, 4),
    ):
        st = style_by_name(styles, style_name)
        st.font.name = head_west
        st._element.rPr.rFonts.set(qn("w:eastAsia"), head_east)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    first_heading_seen = False
    in_references = False
    in_related = False
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Heading 2":
            paragraph.style = style_by_name(styles, "Heading 1")
        elif paragraph.style.name == "Heading 3":
            paragraph.style = style_by_name(styles, "Heading 2")

        if paragraph.style.name.startswith("Heading"):
            if not first_heading_seen and paragraph.text.strip() in {"Abstract", "초록"}:
                paragraph.paragraph_format.page_break_before = True
                first_heading_seen = True
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
            if paragraph.text.strip() in {"References", "참고문헌"}:
                in_references = True
                in_related = False
            elif paragraph.text.strip() in {"Related WRRA Records", "연관 WRRA 기록"}:
                in_references = False
                in_related = True
        elif not first_heading_seen:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(5)

        xml = paragraph._p.xml
        if "m:oMath" in xml and not paragraph.text.strip():
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(7)
            paragraph.paragraph_format.keep_together = True

        for run in paragraph.runs:
            if paragraph.style.name == "Title":
                set_font(run, head_west, head_east, 25, True)
            elif paragraph.style.name == "Subtitle":
                set_font(run, head_west, head_east, 14, False)
            elif paragraph.style.name.startswith("Heading"):
                set_font(run, head_west, head_east, None, None)
            else:
                set_font(run, body_west, body_east, None, None)
                if in_references:
                    run.font.size = Pt(9.1)
                elif in_related:
                    run.font.size = Pt(9.6)
        if in_references and not paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.paragraph_format.space_after = Pt(1.5)
        elif in_related and not paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.line_spacing = 1.02
            paragraph.paragraph_format.space_after = Pt(1.5)

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        table.rows[0].height = None
        set_repeat_table_header(table.rows[0])
        for r_idx, row in enumerate(table.rows):
            prevent_row_split(row)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                set_cell_borders(cell)
                if r_idx == 0:
                    set_cell_shading(cell, "23364D")
                elif r_idx % 2 == 0:
                    set_cell_shading(cell, "F3F6F8")
                else:
                    set_cell_shading(cell, "FFFFFF")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(1)
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.04
                    paragraph.paragraph_format.widow_control = True
                    for run in paragraph.runs:
                        set_font(
                            run,
                            head_west if r_idx == 0 else body_west,
                            head_east if r_idx == 0 else body_east,
                            8.2 if len(table.columns) >= 4 else 8.8,
                            True if r_idx == 0 else None,
                            "FFFFFF" if r_idx == 0 else "000000",
                        )

    for section in doc.sections:
        footer = section.footer
        if not footer.paragraphs:
            footer.add_paragraph()
        p = footer.paragraphs[0]
        p.clear()
        add_page_number(p)
        for run in p.runs:
            set_font(run, head_west, head_east, 8.5, False, "666666")

    props = doc.core_properties
    props.author = "Wonsik Choi"
    props.last_modified_by = "Wonsik Choi"
    props.subject = "WRRA hybrid expansion twist cosmology and hidden gravitational state"
    props.keywords = "WRRA, finite boundaryless universe, holonomy, twist stress, cosmological redshift"
    props.comments = "Version 1.0  28 September 2026"

    doc.save(path)


if __name__ == "__main__":
    style_doc(PAPER / "WRRA_Finite_Boundaryless_EN_v1.0.docx", korean=False)
    style_doc(PAPER / "WRRA_Finite_Boundaryless_KO_v1.0.docx", korean=True)
